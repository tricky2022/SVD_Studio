"""Import tabular device descriptions into the domain model.

The importer is intentionally deterministic. An AI or user may produce a mapping
plan, but only this module converts approved rows into domain objects.
"""
from __future__ import annotations

import csv
import json
import re
from collections.abc import Iterable, Mapping
from dataclasses import dataclass, field
from pathlib import Path

from svdstudio.domain.model import (
    AddressBlock,
    EnumeratedValues,
    EnumValue,
    SvdDevice,
    SvdField,
    SvdPeripheral,
    SvdRegister,
)

ALIASES = {
    "peripheral": ("peripheral", "periph", "module", "block"),
    "base_address": ("base_address", "base", "peripheral_base", "baseaddress"),
    "register": ("register", "reg", "register_name"),
    "address_offset": ("address_offset", "offset", "reg_offset", "addressoffset"),
    "register_size": ("register_size", "reg_size", "size", "width"),
    "access": ("access", "permission", "permissions"),
    "reset_value": ("reset_value", "reset", "resetvalue"),
    "reset_mask": ("reset_mask", "resetmask"),
    "description": ("description", "desc", "comment", "comments"),
    "field": ("field", "field_name", "bitfield"),
    "bit_offset": ("bit_offset", "lsb", "bitoffset", "bit_start"),
    "bit_width": ("bit_width", "width_bits", "bitwidth"),
    "msb": ("msb", "bit_end"),
    "enum_group": ("enum_group", "enumerated_values", "enum_set"),
    "enum_name": ("enum_name", "enum_value_name", "enum"),
    "enum_value": ("enum_value", "value", "enumerated_value"),
    "enum_description": ("enum_description", "enum_desc"),
}


@dataclass
class ImportIssue:
    severity: str
    row: int
    message: str
    field: str = ""


@dataclass
class ImportResult:
    device: SvdDevice | None
    issues: list[ImportIssue] = field(default_factory=list)
    rows_read: int = 0
    peripherals_created: int = 0
    registers_created: int = 0
    fields_created: int = 0

    @property
    def valid(self) -> bool:
        return self.device is not None and not any(i.severity == "error" for i in self.issues)


def _key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value).strip().lower()).strip("_")


def _parse_int(value: object, default: int | None = None) -> int | None:
    if value is None or str(value).strip() == "":
        return default
    text = str(value).strip().replace("_", "")
    try:
        return int(text, 0)
    except ValueError:
        try:
            return int(text, 16) if re.fullmatch(r"[0-9a-fA-F]+", text) else default
        except ValueError:
            return default


def _normalized(row: Mapping[str, object], mapping: Mapping[str, str]) -> dict[str, str]:
    source = {_key(k): "" if v is None else str(v).strip() for k, v in row.items()}
    result = {}
    for canonical, aliases in ALIASES.items():
        source_column = mapping.get(canonical, "")
        if source_column:
            result[canonical] = source.get(_key(source_column), "")
            continue
        for alias in aliases:
            if alias in source:
                result[canonical] = source[alias]
                break
        else:
            result[canonical] = ""
    return result


def _default_mapping(headers: Iterable[str]) -> dict[str, str]:
    available = {_key(header): str(header) for header in headers}
    mapping = {}
    for canonical, aliases in ALIASES.items():
        for alias in aliases:
            if alias in available:
                mapping[canonical] = available[alias]
                break
    return mapping


def load_rows(path: str, sheet: str = "") -> tuple[list[dict[str, object]], list[str]]:
    """Read CSV/TSV or XLSX without importing any GUI code."""
    suffix = Path(path).suffix.lower()
    if suffix in {".csv", ".tsv"}:
        with open(path, newline="", encoding="utf-8-sig") as handle:
            dialect = "excel-tab" if suffix == ".tsv" else "excel"
            rows = list(csv.DictReader(handle, dialect=dialect))
        return rows, list(rows[0].keys()) if rows else []
    if suffix in {".xlsx", ".xlsm"}:
        try:
            import openpyxl
        except ImportError as exc:
            raise RuntimeError("Excel support requires: pip install -e '.[excel]'") from exc
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook[sheet] if sheet else workbook.active
        values = worksheet.iter_rows(values_only=True)
        headers = [str(value or "") for value in next(values, ())]
        rows = [dict(zip(headers, values)) for values in values]
        return rows, headers
    raise ValueError(f"Unsupported tabular file: {suffix}; use .csv, .tsv, .xlsx or .xlsm")


def import_rows(
    rows: Iterable[Mapping[str, object]],
    device_name: str,
    mapping: Mapping[str, str] | None = None,
    vendor: str = "",
) -> ImportResult:
    rows = list(rows)
    result = ImportResult(device=SvdDevice(name=device_name, vendor=vendor))
    if not rows:
        result.issues.append(ImportIssue("error", 0, "The table contains no data rows"))
        return result
    mapping = dict(mapping or _default_mapping(rows[0].keys()))
    peripherals: dict[str, SvdPeripheral] = {}
    registers: dict[tuple[str, str], SvdRegister] = {}
    fields: dict[tuple[str, str, str], SvdField] = {}
    enums: dict[tuple[str, str, str, str], EnumeratedValues] = {}

    for row_number, raw in enumerate(rows, start=2):
        result.rows_read += 1
        row = _normalized(raw, mapping)
        peripheral_name = row["peripheral"]
        register_name = row["register"]
        if not peripheral_name or not register_name:
            result.issues.append(ImportIssue("error", row_number, "Peripheral and register are required"))
            continue
        peripheral = peripherals.get(peripheral_name)
        if peripheral is None:
            peripheral = SvdPeripheral(
                name=peripheral_name,
                base_address=_parse_int(row["base_address"], 0) or 0,
                description=row["description"],
            )
            peripherals[peripheral_name] = peripheral
            result.device.peripherals.append(peripheral)
        elif row["base_address"] and peripheral.base_address != (_parse_int(row["base_address"], 0) or 0):
            result.issues.append(ImportIssue("warning", row_number, "Conflicting peripheral base address", "base_address"))

        register_key = (peripheral_name, register_name)
        register = registers.get(register_key)
        if register is None:
            register = SvdRegister(
                name=register_name,
                address_offset=_parse_int(row["address_offset"], 0) or 0,
                size=_parse_int(row["register_size"], result.device.width) or result.device.width,
                access=row["access"],
                reset_value=_parse_int(row["reset_value"], 0) or 0,
                reset_mask=_parse_int(row["reset_mask"], 0xFFFFFFFF) or 0,
                description=row["description"],
            )
            registers[register_key] = register
            peripheral.registers.append(register)
            result.registers_created += 1
        if not row["field"]:
            continue
        field_key = (peripheral_name, register_name, row["field"])
        field = fields.get(field_key)
        bit_offset = _parse_int(row["bit_offset"])
        bit_width = _parse_int(row["bit_width"])
        if bit_width is None and row["field"]:
            bit_width = _parse_int(row["register_size"])
        msb = _parse_int(row["msb"])
        if bit_offset is None and msb is not None and bit_width is not None:
            bit_offset = msb - bit_width + 1
        if bit_width is None and msb is not None and bit_offset is not None:
            bit_width = msb - bit_offset + 1
        if bit_offset is None or bit_width is None:
            result.issues.append(ImportIssue("error", row_number, "Field requires bit offset/width or MSB/width", "bit_offset"))
            continue
        if field is None:
            field = SvdField(
                name=row["field"], description=row["description"], bit_offset=bit_offset,
                bit_width=bit_width, access=row["access"], reset_value=_parse_int(row["reset_value"]),
            )
            fields[field_key] = field
            register.fields.append(field)
            result.fields_created += 1
        enum_name = row["enum_name"]
        if not enum_name:
            continue
        enum_group = row["enum_group"] or field.name
        enum_key = (*field_key, enum_group)
        enum_values = enums.get(enum_key)
        if enum_values is None:
            enum_values = EnumeratedValues(name=enum_group)
            enums[enum_key] = enum_values
            field.enumerated_values.append(enum_values)
        enum_value = _parse_int(row["enum_value"])
        if enum_value is None:
            result.issues.append(ImportIssue("error", row_number, "Enumeration requires a numeric value", "enum_value"))
        else:
            enum_values.values.append(EnumValue(name=enum_name, description=row["enum_description"], value=enum_value))

    result.peripherals_created = len(peripherals)
    for peripheral in result.device.peripherals:
        max_end = max((r.address_offset + max(1, r.size // 8) for r in peripheral.registers), default=0)
        if max_end:
            peripheral.address_blocks.append(AddressBlock(offset=0, size=max_end))
    return result


def import_file(path: str, device_name: str, mapping_file: str = "", sheet: str = "", vendor: str = "") -> ImportResult:
    rows, headers = load_rows(path, sheet)
    mapping = json.loads(Path(mapping_file).read_text(encoding="utf-8")) if mapping_file else _default_mapping(headers)
    return import_rows(rows, device_name, mapping, vendor)


def normalize_header(value: object) -> str:
    return _key(value)


def suggest_mapping(headers: Iterable[str]) -> dict[str, str]:
    return _default_mapping(headers)


def list_excel_sheets(path: str) -> list[str]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("Excel support requires: pip install -e '.[excel]'") from exc
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return list(workbook.sheetnames)


def preview_table(path: str, sheet: str | None = None, limit: int = 30) -> tuple[list[str], list[dict]]:
    rows, headers = load_rows(path, sheet or "")
    return headers, rows[:limit]
